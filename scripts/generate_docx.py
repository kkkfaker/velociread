#!/usr/bin/env python3
"""
paper-anatomist Word Report Generator
将精读笔记生成格式化的 .docx 文件，存入 Zotero 文献目录。

Usage:
    python generate_docx.py <notes.json> --output <output.docx>
"""

import argparse, json, sys
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Inches, Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
except ImportError:
    print("需要 python-docx: pip install python-docx")
    sys.exit(1)


def build_docx(data: dict, output_path: Path) -> None:
    doc = Document()

    # Styles
    style = doc.styles['Normal']
    style.font.size = Pt(11)
    style.font.name = 'Calibri'

    # --- Title ---
    title = doc.add_heading(data.get('title', '精读笔记'), level=1)
    doc.add_paragraph(
        f"期刊: {data.get('journal', '')} | 年份: {data.get('year', '')} | DOI: {data.get('doi', '')}\n"
        f"作者: {data.get('authors', '')}\n"
        f"日期: {data.get('date', '')} | 深度: {data.get('depth', '')} | 翻译: {data.get('engine', '')}"
    )

    # --- Quick Summary ---
    doc.add_heading('速览卡片', level=2)
    summary = data.get('summary', {})
    table = doc.add_table(rows=5, cols=2, style='Light Grid Accent 1')
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    items = [('核心问题', 'question'), ('核心方法', 'method'),
             ('关键结果', 'result'), ('意义', 'significance'), ('适合引用于', 'cite_in')]
    for i, (label, key) in enumerate(items):
        table.cell(i, 0).text = f'**{label}**'
        table.cell(i, 1).text = summary.get(key, '')

    # --- Terminology ---
    terms = data.get('terminology', [])
    if terms:
        doc.add_heading('术语表', level=2)
        t = doc.add_table(rows=len(terms) + 1, cols=4, style='Light Grid Accent 1')
        for j, h in enumerate(['英文术语', '中文翻译', '定义', '类别']):
            t.cell(0, j).text = h
        for i, term in enumerate(terms):
            t.cell(i + 1, 0).text = term.get('en', '')
            t.cell(i + 1, 1).text = term.get('zh', '')
            t.cell(i + 1, 2).text = term.get('def', '')
            t.cell(i + 1, 3).text = term.get('category', '')

    # --- Sections ---
    for section in data.get('sections', []):
        doc.add_heading(section.get('heading', ''), level=2)

        for para in section.get('paragraphs', []):
            doc.add_paragraph(para.get('text', ''), style='List Bullet')

        for highlight in section.get('highlights', []):
            p = doc.add_paragraph()
            run = p.add_run(f"🔴 {highlight.get('text', '')}")
            run.font.color.rgb = RGBColor(0xCC, 0x33, 0x33)
            if highlight.get('note'):
                p.add_run(f"\n    💡 {highlight['note']}").font.italic = True

    # --- Figures ---
    figures = data.get('figures', [])
    if figures:
        doc.add_heading('图表解读', level=2)
        for fig in figures:
            doc.add_heading(f"Figure {fig.get('num', '')} | 图{fig.get('num', '')}", level=3)
            img_path = fig.get('image_path', '')
            if img_path and Path(img_path).exists():
                doc.add_picture(img_path, width=Inches(5))
            p = doc.add_paragraph()
            p.add_run(f"基本信息: {fig.get('info', '')}\n").bold = True
            p.add_run(f"主要结论: {fig.get('conclusion', '')}\n")
            p.add_run(f"原文结论句: 🟣 {fig.get('source', '')}").font.color.rgb = RGBColor(0x66, 0x33, 0xCC)

    # --- Gems (Deep mode) ---
    gems = data.get('gems', [])
    if gems:
        doc.add_heading('精辟句子收藏', level=2)
        for g in gems:
            doc.add_paragraph(f'"{g["text"]}" — {g["note"]}', style='Quote')

    # --- Review Cards (Deep mode) ---
    cards = data.get('review_cards', [])
    if cards:
        doc.add_heading('复习卡片', level=2)
        for c in cards:
            doc.add_paragraph(f"Q: {c['q']}", style='List Bullet')
            doc.add_paragraph(f"A: {c['a']}")

    # --- Log ---
    doc.add_heading('阅读日志', level=2)
    log = data.get('log', {})
    for stage, detail in log.items():
        doc.add_paragraph(f"{stage}: {detail}")

    doc.save(str(output_path))


def main():
    parser = argparse.ArgumentParser(description="Generate formatted Word report from reading notes JSON")
    parser.add_argument("notes_json", help="Notes JSON file path")
    parser.add_argument("--output", "-o", required=True, help="Output .docx path")
    args = parser.parse_args()

    with open(args.notes_json, 'r', encoding='utf-8') as f:
        data = json.load(f)

    output = Path(args.output)
    build_docx(data, output)
    print(f"✅ Word report: {output} ({output.stat().st_size / 1024:.0f} KB)")


if __name__ == "__main__":
    main()
